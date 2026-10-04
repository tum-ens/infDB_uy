"""INE: census cartography 2023 (GeoPackage) and the official variable dictionary.

Cartography layers are reprojected to the project CRS (source: EPSG:4326) and kept when
they intersect the scope. The dictionary is parsed into tidy tables; the xlsx stays in raw/.
INE's generic codes (7777 no corresponde, 8888 no relevado, 9898 ignorado, 99/9999 no
recuerda) are listed as codes, never turned into missing values.
"""

from __future__ import annotations

import logging
import zipfile
from pathlib import Path

import geopandas as gpd
import openpyxl
import pandas as pd

from ..core import Context, write_json
from ..scope import annotate_scope, load_scope

log = logging.getLogger(__name__)


def cartografia(ctx: Context) -> None:
    url = ctx.source("ine_cartografia")["url"]
    zpath = ctx.download(url, ctx.raw("ine_cartografia", Path(url.replace("%20", " ")).name), "ine_cartografia")
    scope = load_scope(ctx)
    out = ctx.prepared("ine_cartografia.gpkg")
    summary = {}
    with zipfile.ZipFile(zpath) as zf:
        members = [n for n in zf.namelist() if n.lower().endswith(".gpkg")]
        for pdf in (n for n in zf.namelist() if n.lower().endswith(".pdf")):
            target = ctx.raw("ine_cartografia", "metadata", Path(pdf).name)
            if not target.exists():
                target.write_bytes(zf.read(pdf))
    for m in members:
        layer = Path(m).stem
        gdf = gpd.read_file(f"/vsizip/{zpath}/{m}", engine="pyogrio")
        src_crs = gdf.crs.to_string() if gdf.crs else None
        gdf = annotate_scope(gdf.to_crs(ctx.crs), scope)
        gdf.to_file(out, layer=layer, driver="GPKG")
        summary[layer] = {"source_crs": src_crs, "features_intersecting_scope": len(gdf),
                          "by_scope_relation": gdf["scope_relation"].value_counts().to_dict() if len(gdf) else {}}
        log.info("INE %s: %d features intersect scope", layer, len(gdf))
    write_json(ctx.reports("ine_cartografia_summary.json"), summary)


def parse_dictionary(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return (variables, categories, generic_codes) from INE's 'Diccionario de variables'."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    variables, categories = [], []
    for ws in wb.worksheets:
        title = ws.title.strip().upper()
        if title not in {"VIVIENDA", "HOGAR", "PERSONA"}:
            continue
        section, current = None, None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 5:
                continue
            a, b, c, d, e = (list(row) + [None] * 5)[:5]
            a = a.strip() if isinstance(a, str) else a
            if b is not None and str(b).strip():
                current = {"table": title, "section": section, "label": a, "name": str(b).strip(),
                           "description": None, "universe": e, "row": i + 1}
                variables.append(current)
                if c is not None:
                    categories.append({"table": title, "name": current["name"], "code": str(c).strip(), "label": d})
                elif d is not None:
                    current["description"] = d
            elif c is not None and current is not None:
                categories.append({"table": title, "name": current["name"], "code": str(c).strip(), "label": d})
            elif a is not None and b is None and c is None:
                section = a
            elif d is not None and current is not None:
                current["description"] = ((current["description"] or "") + " " + str(d)).strip()
    generic = []
    if "CÓDIGOS GENÉRICOS" in wb.sheetnames:
        for row in list(wb["CÓDIGOS GENÉRICOS"].iter_rows(values_only=True))[2:]:
            if row and row[0] is not None:
                generic.append({"code": str(row[0]).strip(), "label": row[1]})
    return pd.DataFrame(variables), pd.DataFrame(categories), pd.DataFrame(generic)


def diccionario(ctx: Context) -> None:
    url = ctx.source("ine_diccionario")["url"]
    path = ctx.download(url, ctx.raw("ine_diccionario", Path(url.replace("%20", " ")).name), "ine_diccionario")
    variables, categories, generic = parse_dictionary(path)
    variables.to_parquet(ctx.prepared("ine", "diccionario_variables.parquet"), index=False)
    categories.to_parquet(ctx.prepared("ine", "diccionario_categorias.parquet"), index=False)
    generic.to_parquet(ctx.prepared("ine", "codigos_genericos.parquet"), index=False)
    log.info("INE dictionary: %d variables, %d categories", len(variables), len(categories))


def run(ctx: Context) -> None:
    if ctx.enabled("ine_cartografia"):
        cartografia(ctx)
    if ctx.enabled("ine_diccionario"):
        diccionario(ctx)
